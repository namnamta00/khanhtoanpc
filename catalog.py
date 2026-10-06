"""Đọc báo giá, không thực thi công thức Excel."""
import csv
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path


def normalized(value):
    value = unicodedata.normalize('NFD', str(value).lower().replace('đ', 'd'))
    return re.sub(r'[^a-z0-9]', '', ''.join(c for c in value if not unicodedata.combining(c)))


@dataclass
class Product:
    name: str
    price: str = ''
    description: str = ''
    code: str = ''


ALIASES = {
    'name': ['tên sản phẩm', 'sản phẩm', 'tên hàng', 'tên hàng hóa', 'product', 'name'],
    'price': ['giá', 'đơn giá', 'giá bán', 'giá niêm yết', 'price'],
    'description': ['mô tả', 'cấu hình', 'thông số', 'description'],
    'code': ['mã', 'mã sản phẩm', 'mã hàng', 'sku', 'code'],
}


def read_catalog(path):
    path = Path(path)
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()
    elif path.suffix.lower() == '.csv':
        raw = path.read_bytes()
        try:
            content = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            content = raw.decode('cp1258')
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(content.splitlines(), dialect))
    else:
        raise ValueError('Hãy chọn file Excel .xlsx hoặc CSV. File .xls cần lưu lại thành .xlsx.')
    aliases = {key: {normalized(x) for x in names} for key, names in ALIASES.items()}
    for header_index, row in enumerate(rows[:30]):
        mapping = {key: next((i for i, val in enumerate(row) if normalized(val) in names), None)
                   for key, names in aliases.items()}
        if mapping['name'] is not None:
            break
    else:
        raise ValueError('Không tìm thấy cột Tên sản phẩm trong 30 dòng đầu. Xem file báo giá mẫu.')
    products = []
    for row in rows[header_index + 1:]:
        values = {}
        for key, index in mapping.items():
            val = row[index] if index is not None and index < len(row) else None
            if key == 'price' and isinstance(val, (float, int)):
                val = f'{val:,.0f}'.replace(',', '.') + ' đ'
            values[key] = str(val).strip() if val is not None else ''
        if values['name']:
            products.append(Product(**values))
    if not products:
        raise ValueError('File báo giá chưa có sản phẩm.')
    return products


def hashtags(product):
    words = re.findall(r'[A-Za-z0-9]+', unicodedata.normalize('NFKD', product.name.replace('đ', 'd')).encode('ascii', 'ignore').decode())
    tags = ['khanhtoan'] + [word.lower() for word in words[:8]]
    return ' '.join('#' + x for x in dict.fromkeys(tags) if x)


STORE_FOOTER = '''----------
★ KHÁNH TOÀN COMPUTER ★ Uy Tín - Chất Lượng - Chuyên Nghiệp
🌐 Website: https://www.khanhtoan.com/
📞 Hotline: 085.597.6363
🏢 Cơ Sở 1: số 173 Quang Trung, Phường Nam Định, tỉnh Ninh Bình, Việt Nam.
🏢 Cơ Sở 2: số 24 Nguyễn Lương Bằng, Phường Hoa Lư, tỉnh Ninh Bình, Việt Nam.'''


def finish_post(content, product):
    """Append the supplied store details once, keeping hashtags at the bottom."""
    content = clean_post(content)
    # The model writes only the product section; recover if it repeats the footer.
    content = re.split(r'(?m)^\s*(?:-{3,}|★\s*KHÁNH TOÀN COMPUTER)', content, maxsplit=1)[0]
    tags = re.findall(r'(?<!\w)#[\w]+', content)
    content = re.sub(r'(?<!\w)#[\w]+', '', content).strip()
    if not content:
        raise RuntimeError('AI chưa trả về nội dung sản phẩm. Hãy thử tạo lại.')
    combined = ['#khanhtoan'] + tags + hashtags(product).split()
    unique = {tag.casefold(): tag for tag in combined}
    return content + '\n' + STORE_FOOTER + '\n' + ' '.join(unique.values())


def clean_post(text):
    """Remove recognizable assistant preambles only at the start of a post."""
    lines = text.strip().splitlines()
    while lines:
        first = lines[0].strip()
        plain = first.strip('*_ ').casefold()
        intro = re.match(
            r'^(?:(?:tuyệt vời|chắc chắn rồi|được rồi|vâng|dĩ nhiên)[!,. :–—-]*\s*)?'
            r'(?:dưới đây là|đây là)\s+(?:bài viết|bài đăng|nội dung bài|nội dung bán hàng)', plain)
        acknowledgement = plain in {'tuyệt vời!', 'chắc chắn rồi!', 'được rồi!', 'vâng!'}
        if not first or intro or acknowledgement or plain in {'---', '```', '```text', '```plaintext'}:
            lines.pop(0)
        else:
            break
    if lines and lines[-1].strip() == '```':
        lines.pop()
    return '\n'.join(lines).strip()


def draft(product, extra):
    points = [line.strip().lstrip('✅•- ').strip()
              for line in (product.description + '\n' + extra).splitlines() if line.strip()]
    parts = [f'✨ {product.name.upper()}'] + ['✅ ' + point for point in points if point]
    if product.price:
        parts.append(f'💰 Giá: {product.price}')
    return finish_post('\n'.join(parts), product)
